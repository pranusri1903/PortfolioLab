import io
import json
from datetime import date, datetime, timedelta
from decimal import Decimal as D

import pytest
from openpyxl import Workbook

from app.api.market import get_provider
from app.main import app
from app.services.live_data import CompositeProvider, SymbolHit
from tests.conftest import auth

BASE = "/api/v1/portfolios"
H = auth("imp")
CANON = ["trade_date", "type", "symbol", "quantity", "price", "fee", "cash_amount", "notes"]
FUND = "Test Flexi Cap Fund - Direct Growth"


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def new_portfolio(client, currency="USD"):
    r = client.post(BASE, json={"name": "P", "currency": currency}, headers=H)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def upload(client, pid, content: bytes, name: str, path="transactions/import", commit=False, **form):
    return client.post(
        f"{BASE}/{pid}/{path}?commit={str(commit).lower()}",
        headers=H,
        files={"file": (name, io.BytesIO(content))},
        data=form,
    )


class FakeFunds:
    def lookup(self, code):
        return SymbolHit("MF-1", FUND, "MUTUAL_FUND", "INR", "AMFI") if code == "1" else None


class FundProvider(CompositeProvider):
    source = "fake"

    def __init__(self):
        super().__init__(stocks=None, funds=FakeFunds())  # type: ignore[arg-type]

    def search(self, q, currency=None):
        ok = currency in (None, "INR") and q.lower().startswith("test flexi")
        return [SymbolHit("MF-1", FUND, "MUTUAL_FUND", "INR", "AMFI")] if ok else []

    def source_for(self, a):
        return "fake"

    def daily_closes(self, asset, start, end):
        d = max(start, date.today() - timedelta(days=400))
        while d <= end:
            if d.weekday() < 5:
                yield d, D("50") + D(d.toordinal() % 9)
            d += timedelta(days=1)


@pytest.fixture
def funds(client):
    app.dependency_overrides[get_provider] = lambda: FundProvider()


def test_excel_with_native_dates_and_numbers(client):
    pid = new_portfolio(client)
    book = xlsx(
        [
            CANON,
            [datetime(2025, 1, 10), "DEPOSIT", None, None, None, None, 1000, "cash"],
            [datetime(2025, 1, 13), "BUY", "ACME", 10, 50.25, 1, None, "xl"],
        ]
    )
    r = upload(client, pid, book, "history.xlsx").json()
    assert r["committed"] is False and (r["valid"], r["invalid"]) == (2, 0)
    done = upload(client, pid, book, "history.xlsx", commit=True).json()
    assert done["result"]["imported"] == 2
    h = client.get(f"{BASE}/{pid}/holdings", headers=H).json()["holdings"][0]
    assert h["symbol"] == "ACME" and D(h["cost_basis"]) == D("503.50")


def test_excel_title_rows_above_the_header_are_skipped(client):
    pid = new_portfolio(client)
    book = xlsx([["Tradebook FY25"], [], CANON, ["2025-01-10", "DEPOSIT", None, None, None, None, 500, None]])
    assert upload(client, pid, book, "t.xlsx").json()["valid"] == 1


def test_bad_files_get_clear_errors(client):
    pid = new_portfolio(client)
    assert upload(client, pid, b"junk", "x.xls").json()["error"]["code"] == "invalid_file"
    assert (
        upload(client, pid, b"PK\x03\x04not really a zip", "x.xlsx").json()["error"]["code"] == "invalid_file"
    )
    assert (
        upload(client, pid, xlsx([["a", "b", "c"], [1, 2, 3]]), "x.xlsx").json()["error"]["code"]
        == "invalid_csv_header"
    )


def test_asset_class_scope_keeps_stocks_and_funds_apart(client):
    pid = new_portfolio(client)
    book = xlsx(
        [
            CANON,
            ["2025-01-10", "BUY", "ACME", 1, 50, 0, None, None],
            ["2025-01-10", "BUY", "GRWF", 1, 40, 0, None, None],
        ]
    )
    stocks = upload(client, pid, book, "m.xlsx", asset_class="stock", auto_fund="true").json()
    assert [r["status"] for r in stocks["rows"]] == ["valid", "invalid"]
    assert "import it under mutual funds" in stocks["rows"][1]["reasons"][0]
    funds_ = upload(client, pid, book, "m.xlsx", asset_class="fund", auto_fund="true").json()
    assert [r["status"] for r in funds_["rows"]] == ["invalid", "valid"]


def test_files_without_deposits_need_auto_fund(client):
    pid = new_portfolio(client)
    book = xlsx([CANON, ["2025-01-10", "BUY", "ACME", 10, 50, 0, None, None]])
    plain = upload(client, pid, book, "b.xlsx").json()
    assert plain["invalid"] == 1 and "negative" in plain["rows"][0]["reasons"][0]
    funded = upload(client, pid, book, "b.xlsx", commit=True, auto_fund="true").json()
    assert funded["result"]["imported"] == 1
    s = client.get(f"{BASE}/{pid}/summary", headers=H).json()
    assert D(s["cash_balance"]) == 0 and D(s["net_contributions"]) == 500


def test_mutual_fund_statement_resolves_scheme_names_and_derives_units(client, funds):
    pid = new_portfolio(client, "INR")
    book = xlsx(
        [
            ["Consolidated Account Statement"],
            [],
            ["Date", "Scheme Name", "Transaction", "Amount (INR)", "NAV", "Units"],
            [datetime(2025, 1, 6), FUND, "Purchase", "10,000.00", 50.0, None],
            [datetime(2025, 2, 5), FUND, "Systematic Investment", 5000, 52.5, None],
            [datetime(2025, 3, 5), FUND, "Redemption", 4000, 55, None],
            [datetime(2025, 3, 6), "Some Unknown Fund", "Purchase", 100, 10, None],
        ]
    )
    mapping = json.dumps(
        {
            "trade_date": "Date",
            "type": "Transaction",
            "symbol": "Scheme Name",
            "cash_amount": "Amount (INR)",
            "price": "NAV",
            "quantity": "Units",
        }
    )
    kw = dict(mapping=mapping, date_format="%d/%m/%Y", asset_class="fund", auto_fund="true")
    pre = upload(client, pid, book, "cas.xlsx", **kw).json()
    assert [r["status"] for r in pre["rows"]] == ["valid", "valid", "valid", "invalid"]
    assert pre["rows"][0]["resolved_symbol"] == "MF-1" and pre["rows"][0]["will_add_asset"]
    assert "scheme name or its AMFI code" in pre["rows"][3]["reasons"][0]
    assert client.get(f"{BASE}/{pid}/holdings", headers=H).json()["holdings"] == []  # preview wrote nothing
    done = upload(client, pid, book, "cas.xlsx", commit=True, **kw).json()
    assert done["result"] == {"imported": 3, "skipped": 0, "rejected": 1}
    h = client.get(f"{BASE}/{pid}/holdings", headers=H).json()["holdings"][0]
    assert h["symbol"] == "MF-1" and h["asset_type"] == "MUTUAL_FUND"
    assert abs(D(h["quantity"]) - (D("200") + D("95.23809524") - D("72.72727273"))) < D("0.0001")
    s = client.get(f"{BASE}/{pid}/summary", headers=H).json()
    assert D(s["cash_balance"]) >= 0 and D("15000") <= D(s["net_contributions"]) < D("15000.01")


def test_fund_resolved_by_amfi_code(client, funds):
    pid = new_portfolio(client, "INR")
    book = xlsx([["Date", "Code", "Type", "Amount", "NAV"], ["2025-01-06", 1, "Purchase", 1000, 50]])
    mapping = json.dumps(
        {"trade_date": "Date", "type": "Type", "symbol": "Code", "cash_amount": "Amount", "price": "NAV"}
    )
    r = upload(client, pid, book, "c.xlsx", mapping=mapping, asset_class="fund", auto_fund="true").json()
    assert r["rows"][0]["status"] == "valid" and r["rows"][0]["resolved_symbol"] == "MF-1"


def test_holdings_file_builds_positions_with_funding_and_is_idempotent(client):
    pid = new_portfolio(client)
    book = xlsx(
        [
            ["Symbol", "Qty", "Avg Price", "Buy Date"],
            ["ACME", 10, 50, datetime(2025, 1, 10)],
            ["CRST", 5, 100, None],
        ]
    )
    form = dict(default_date="2025-02-03", cash="250")
    pre = upload(client, pid, book, "h.xlsx", path="holdings/import", **form).json()
    assert (pre["valid"], pre["invalid"]) == (2, 0)
    done = upload(client, pid, book, "h.xlsx", path="holdings/import", commit=True, **form).json()
    assert done["result"]["imported"] == 2
    s = client.get(f"{BASE}/{pid}/summary", headers=H).json()
    assert D(s["cash_balance"]) == 250 and D(s["net_contributions"]) == 1250
    assert {h["symbol"] for h in client.get(f"{BASE}/{pid}/holdings", headers=H).json()["holdings"]} == {
        "ACME",
        "CRST",
    }
    assert (
        client.get(f"{BASE}/{pid}/transactions", headers=H).json()["total"] == 5
    )  # 2 deposits + 2 buys + opening cash
    again = upload(client, pid, book, "h.xlsx", path="holdings/import", commit=True, **form).json()
    assert again["result"] == {"imported": 0, "skipped": 2, "rejected": 0}
    assert client.get(f"{BASE}/{pid}/transactions", headers=H).json()["total"] == 5  # cash not added twice


def test_holdings_file_can_use_invested_amount_and_csv(client):
    pid = new_portfolio(client)
    csv = "scheme,units,invested amount,date\nGRWF,25,1000,2025-01-10\n"
    r = upload(
        client,
        pid,
        csv.encode(),
        "h.csv",
        path="holdings/import",
        commit=True,
        default_date="2025-02-03",
        asset_class="fund",
    ).json()
    assert r["result"]["imported"] == 1
    h = client.get(f"{BASE}/{pid}/holdings", headers=H).json()["holdings"][0]
    assert D(h["average_cost"]) == 40 and D(h["quantity"]) == 25


def test_holdings_file_validation(client):
    pid = new_portfolio(client)
    ok = b"symbol,quantity,average_price\nACME,1,50\n"
    r = upload(
        client, pid, ok, "h.csv", path="holdings/import", asset_class="fund", default_date="2025-02-03"
    ).json()
    assert r["rows"][0]["status"] == "invalid" and "mutual funds" in r["rows"][0]["reasons"][0]
    inr = upload(
        client,
        pid,
        b"symbol,quantity,average_price\nINFX,1,50\n",
        "h.csv",
        path="holdings/import",
        default_date="2025-02-03",
    ).json()
    assert "trades in INR, not USD" in inr["rows"][0]["reasons"][0]
    future = (date.today() + timedelta(days=3)).isoformat()
    assert upload(client, pid, ok, "h.csv", path="holdings/import", default_date=future).status_code == 422
    no_cols = upload(
        client, pid, b"foo,bar,baz\n1,2,3\n", "h.csv", path="holdings/import", default_date="2025-02-03"
    )
    assert no_cols.json()["error"]["code"] == "invalid_csv_header"


def test_inspect_returns_columns_and_samples(client):
    book = xlsx([["Junk title"], ["Date", "Scheme", "Units"], ["2025-01-01", "A", 1], ["2025-01-02", "B", 2]])
    r = client.post("/api/v1/import/inspect", headers=H, files={"file": ("x.xlsx", io.BytesIO(book))}).json()
    assert (
        r["columns"] == ["Date", "Scheme", "Units"]
        and r["row_count"] == 2
        and r["sample"][1]["Scheme"] == "B"
    )
    assert (
        client.post("/api/v1/import/inspect", files={"file": ("x.csv", io.BytesIO(b"a,b,c\n"))}).status_code
        == 401
    )


# ---- statements with several tables in one sheet (e.g. a Groww-style holdings export) -----------

GROWW_FUNDS = {
    "sbi gold direct plan growth": SymbolHit(
        "MF-119788", "SBI GOLD FUND - Direct Plan - Growth", "MUTUAL_FUND", "INR", "AMFI"
    ),
    "hdfc infrastructure fund direct growth": SymbolHit(
        "MF-118979", "HDFC Infrastructure Fund - Direct Plan - Growth Option", "MUTUAL_FUND", "INR", "AMFI"
    ),
}


class GrowwProvider(FundProvider):
    def search(self, q, currency=None):
        hit = GROWW_FUNDS.get(q.lower())
        return [hit] if hit and currency in (None, "INR") else []


def groww_sheet(extra_rows=()):
    """Synthetic data in the layout of a mutual-fund holdings export: details, summary, then holdings."""
    return xlsx(
        [
            ["Personal Details"],
            ["Name", "Test User"],
            ["Mobile Number", "0000000000"],
            ["PAN", "AAAAA0000A"],
            [],
            [],
            ["HOLDING SUMMARY"],
            [],
            ["Total Investments", "Current Portfolio Value", "Profit/Loss", "Profit/Loss %", "XIRR"],
            [76496.2, 75291.72, -1204.49, "-1.57%", "-4.98%"],
            [],
            [],
            ["HOLDINGS AS ON 2026-10-02"],
            [],
            [
                "Scheme Name",
                "AMC",
                "Category",
                "Sub-category",
                "Folio No.",
                "Source",
                "Units",
                "Invested Value",
                "Current Value",
                "Returns",
                "XIRR",
            ],
            [],
            [
                "SBI Gold Direct Plan Growth",
                "SBI Mutual Fund",
                "Commodities",
                "Gold",
                "111",
                "Groww",
                970.019,
                43497.83,
                43521.07,
                23.2380838,
                "0.18%",
            ],
            [
                "HDFC Infrastructure Fund Direct Growth",
                "HDFC Mutual Fund",
                "Equity",
                "Sectoral",
                "222",
                "Groww",
                634.17,
                32998.37,
                31770.65,
                -1227.72336,
                "-11.8%",
            ],
            *extra_rows,
        ]
    )


def test_inspect_finds_every_table_and_picks_the_holdings(client):
    r = client.post(
        "/api/v1/import/inspect", headers=H, files={"file": ("g.xlsx", io.BytesIO(groww_sheet()))}
    ).json()
    assert [t["title"] for t in r["tables"]] == ["HOLDING SUMMARY", "HOLDINGS AS ON 2026-10-02"]
    assert r["selected"] == 1 and r["columns"][0] == "Scheme Name" and r["row_count"] == 2
    assert "Test User" not in str(r) and "AAAAA0000A" not in str(r)  # personal details never leave the parser
    other = client.post(
        "/api/v1/import/inspect",
        headers=H,
        data={"table_index": "0"},
        files={"file": ("g.xlsx", io.BytesIO(groww_sheet()))},
    ).json()
    assert other["columns"][0] == "Total Investments" and other["row_count"] == 1


def test_holdings_import_reads_the_right_table_with_no_mapping(client, funds):
    app.dependency_overrides[get_provider] = lambda: GrowwProvider()
    pid = new_portfolio(client, "INR")
    kw = dict(path="holdings/import", asset_class="fund", default_date="2025-06-01")
    pre = upload(client, pid, groww_sheet(), "g.xlsx", **kw).json()
    assert [r["status"] for r in pre["rows"]] == ["valid", "valid"], pre
    assert [r["resolved_symbol"] for r in pre["rows"]] == ["MF-119788", "MF-118979"]
    done = upload(client, pid, groww_sheet(), "g.xlsx", commit=True, **kw).json()
    assert done["result"]["imported"] == 2
    held = {h["symbol"]: h for h in client.get(f"{BASE}/{pid}/holdings", headers=H).json()["holdings"]}
    assert D(held["MF-119788"]["quantity"]) == D("970.019")
    assert abs(D(held["MF-119788"]["cost_basis"]) - D("43497.83")) < D(
        "0.01"
    )  # invested value became cost basis
    s = client.get(f"{BASE}/{pid}/summary", headers=H).json()
    assert abs(D(s["net_contributions"]) - D("76496.20")) < D("0.05")


def test_choosing_a_table_and_setting_the_header_row_manually(client):
    pid = new_portfolio(client, "INR")
    mapping = json.dumps({"symbol": "Total Investments", "quantity": "Profit/Loss"})
    first = upload(
        client,
        pid,
        groww_sheet(),
        "g.xlsx",
        path="holdings/import",
        table_index="0",
        mapping=mapping,
        default_date="2025-06-01",
    ).json()
    assert first["total_rows"] == 1  # the summary table, because it was asked for
    by_row = upload(
        client,
        pid,
        groww_sheet(),
        "g.xlsx",
        path="holdings/import",
        header_row="15",
        default_date="2025-06-01",
    ).json()
    assert by_row["total_rows"] == 2
    bad = upload(
        client,
        pid,
        groww_sheet(),
        "g.xlsx",
        path="holdings/import",
        header_row="2",
        default_date="2025-06-01",
    )
    assert (
        bad.status_code == 422 and bad.json()["error"]["code"] == "invalid_csv_header"
    )  # row 2 is "Name, Test User"
    assert (
        upload(
            client,
            pid,
            groww_sheet(),
            "g.xlsx",
            path="holdings/import",
            table_index="9",
            default_date="2025-06-01",
        ).status_code
        == 422
    )


def test_a_total_row_ends_the_table(client):
    pid = new_portfolio(client, "INR")
    sheet = groww_sheet([["Total", "", "", "", "", "", 1604.189, 76496.2, 75291.72, "", ""]])
    r = upload(client, pid, sheet, "g.xlsx", path="holdings/import", default_date="2025-06-01").json()
    assert r["total_rows"] == 2


def test_multi_table_csv_and_second_sheet(client):
    pid = new_portfolio(client)
    csv = "\n".join(
        [
            "Account report",
            "",
            "Totals,Value,Count",
            "1,2,3",
            "",
            "symbol,quantity,average_price,date",
            "ACME,10,50,2025-01-10",
        ]
    )
    r = upload(
        client, pid, csv.encode(), "multi.csv", path="holdings/import", default_date="2025-06-01"
    ).json()
    assert r["total_rows"] == 1 and r["rows"][0]["status"] == "valid"
    wb = Workbook()
    wb.active.title = "About"
    wb.active.append(["Generated by broker"])
    wb.create_sheet("Holdings").append(["symbol", "quantity", "average_price"])
    wb["Holdings"].append(["CRST", 5, 100])
    out = io.BytesIO()
    wb.save(out)
    r2 = upload(
        client, pid, out.getvalue(), "two.xlsx", path="holdings/import", default_date="2025-06-01"
    ).json()
    assert r2["rows"][0]["resolved_symbol"] == "CRST"


def test_fund_names_match_despite_wording_but_never_guess():
    from app.services.resolve import Resolver

    class Stub:
        funds = None

        def __init__(self, hits):
            self.hits, self.queries = hits, []

        def search(self, q, currency=None):
            self.queries.append(q)
            return (
                self.hits if "plan" not in q.lower() else []
            )  # the long wording finds nothing; the short one does

    direct = SymbolHit("MF-1", "Alpha Gold Fund - Direct Plan - Growth", "MUTUAL_FUND", "INR", "AMFI")
    regular = SymbolHit("MF-2", "Alpha Gold Fund - Regular Plan - Growth", "MUTUAL_FUND", "INR", "AMFI")
    stub = Stub([direct, regular])
    got = Resolver({}, stub, "INR").resolve("Alpha Gold Direct Plan Growth")
    assert (
        got and got.symbol == "MF-1" and len(stub.queries) == 2
    )  # retried without filler words, then picked Direct
    assert (
        Resolver({}, Stub([direct, regular]), "INR").resolve("Alpha Gold Plan") is None
    )  # ambiguous: left unresolved
